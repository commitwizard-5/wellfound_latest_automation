import openpyxl
from openpyxl import Workbook
from datetime import datetime
import os

class ExcelHandler:
    def __init__(self, filename='job_applications.xlsx'):
        self.filename = filename
        self.workbook = None
        self.sheet = None
        self.create_or_load_workbook()

    def create_or_load_workbook(self):
        self.file_path = os.path.abspath(self.filename)
        if os.path.exists(self.file_path):
            self.workbook = openpyxl.load_workbook(self.file_path)
            self.sheet = self.workbook.active
            if self.sheet.cell(row=1, column=5).value != "Status":
                self.sheet.cell(row=1, column=5, value="Status")
        else:
            self.workbook = Workbook()
            self.sheet = self.workbook.active
            self.sheet.append(["Date", "Job Link", "Job Description", "Cover Letter", "Status"])

    def add_application(self, link, description, cover_letter, status):
        current_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.sheet.append([current_date, link, description, cover_letter, status])
        

    def save_excel(self):
        self.workbook.save(filename=self.file_path)

